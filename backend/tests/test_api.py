import os
import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from docx import Document
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas
from sqlalchemy import create_engine, inspect, text

os.environ["ANALYSIS_PROVIDER"] = "fake"

from app.main import app, create_app
from app.core.config import settings


client = TestClient(app)


def migrate_database(database_url: str) -> None:
    config = Config("alembic.ini")
    config.attributes["database_url"] = database_url
    command.upgrade(config, "head")


def test_health_check_reports_service_is_available() -> None:
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_capabilities_report_history_is_disabled_by_default() -> None:
    response = client.get("/api/capabilities")

    assert response.status_code == 200
    assert response.json() == {"history_enabled": False}


def test_disabled_app_does_not_expose_history_routes() -> None:
    response = client.get("/api/history")

    assert response.status_code == 404


def test_capabilities_report_history_is_enabled_for_enabled_app() -> None:
    enabled_client = TestClient(create_app(history_enabled=True))

    response = enabled_client.get("/api/capabilities")

    assert response.status_code == 200
    assert response.json() == {"history_enabled": True}


def test_enabled_app_fails_startup_when_database_is_not_migrated(
    tmp_path: Path,
) -> None:
    application = create_app(
        history_enabled=True,
        database_url=f"sqlite:///{tmp_path / 'unmigrated.db'}",
    )

    with pytest.raises(RuntimeError, match="alembic upgrade head"):
        with TestClient(application):
            pass


def test_initial_migration_creates_complete_history_schema(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'schema.db'}"
    migrate_database(database_url)

    columns = {
        column["name"]
        for column in inspect(create_engine(database_url)).get_columns(
            "analysis_history"
        )
    }

    assert columns == {
        "id",
        "job_title",
        "company_name",
        "display_title",
        "resume_text",
        "job_description",
        "resume_source",
        "original_filename",
        "match_score",
        "summary",
        "strengths",
        "matched_requirements",
        "gaps",
        "recommendations",
        "interview_focus",
        "created_at",
    }


def test_successful_text_analysis_is_saved_with_complete_inputs(
    tmp_path: Path,
) -> None:
    database_url = f"sqlite:///{tmp_path / 'history.db'}"
    migrate_database(database_url)

    with TestClient(
        create_app(history_enabled=True, database_url=database_url)
    ) as enabled_client:
        response = enabled_client.post(
            "/api/analyze",
            data={
                "job_title": "  Senior Frontend Engineer  ",
                "company_name": "  Example Corp  ",
                "job_description": "  Build Vue apps.\n\nKeep this spacing.  ",
                "resume_text": "  Frontend engineer\n\nVue specialist.  ",
            },
        )

    assert response.status_code == 200
    assert response.json()["history_status"] == "saved"
    assert isinstance(response.json()["history_id"], int)

    with create_engine(database_url).connect() as connection:
        record = (
            connection.execute(
                text("SELECT * FROM analysis_history WHERE id = :id"),
                {"id": response.json()["history_id"]},
            )
            .mappings()
            .one()
        )

    assert record["job_title"] == "Senior Frontend Engineer"
    assert record["company_name"] == "Example Corp"
    assert record["display_title"] == "Senior Frontend Engineer"
    assert record["job_description"] == "Build Vue apps.\n\nKeep this spacing."
    assert record["resume_text"] == "Frontend engineer\n\nVue specialist."
    assert record["resume_source"] == "text"
    assert record["original_filename"] is None
    assert record["match_score"] == response.json()["match_score"]
    assert record["summary"] == response.json()["summary"]
    created_at = datetime.fromisoformat(record["created_at"])
    assert created_at.utcoffset() == timezone.utc.utcoffset(created_at)
    assert (datetime.now(timezone.utc) - created_at).total_seconds() < 10
    assert json.loads(record["strengths"]) == response.json()["strengths"]
    assert (
        json.loads(record["matched_requirements"])
        == response.json()["matched_requirements"]
    )
    assert json.loads(record["gaps"]) == response.json()["gaps"]
    assert json.loads(record["recommendations"]) == response.json()["recommendations"]
    assert json.loads(record["interview_focus"]) == response.json()["interview_focus"]


def test_successful_file_analysis_saves_text_and_source_metadata_only(
    tmp_path: Path,
) -> None:
    database_url = f"sqlite:///{tmp_path / 'file-history.db'}"
    migrate_database(database_url)
    file_bytes = b"  Frontend engineer\n\nVue specialist.  "

    fallback_title = "A" * 90
    with TestClient(
        create_app(history_enabled=True, database_url=database_url)
    ) as enabled_client:
        response = enabled_client.post(
            "/api/analyze",
            data={"job_description": f"\n\n  {fallback_title}\nSecond line"},
            files={"resume_file": ("resume.txt", file_bytes, "text/plain")},
        )

    assert response.status_code == 200
    with create_engine(database_url).connect() as connection:
        record = (
            connection.execute(text("SELECT * FROM analysis_history")).mappings().one()
        )

    assert record["display_title"] == "A" * 80
    assert record["resume_text"] == "Frontend engineer\n\nVue specialist."
    assert record["resume_source"] == "file"
    assert record["original_filename"] == "resume.txt"


def test_docx_history_preserves_internal_paragraph_formatting(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'docx-history.db'}"
    migrate_database(database_url)
    document_bytes = BytesIO()
    document = Document()
    document.add_paragraph("  Frontend engineer  ")
    document.add_paragraph("")
    document.add_paragraph("Vue specialist")
    document.save(document_bytes)

    with TestClient(
        create_app(history_enabled=True, database_url=database_url)
    ) as enabled_client:
        response = enabled_client.post(
            "/api/analyze",
            data={"job_description": "Build Vue apps."},
            files={
                "resume_file": (
                    "resume.docx",
                    document_bytes.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert response.status_code == 200
    with create_engine(database_url).connect() as connection:
        stored_resume = connection.execute(
            text("SELECT resume_text FROM analysis_history")
        ).scalar_one()
    assert stored_resume == "Frontend engineer  \n\nVue specialist"


def test_validation_failure_creates_no_history_record(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'validation-history.db'}"
    migrate_database(database_url)

    with TestClient(
        create_app(history_enabled=True, database_url=database_url)
    ) as enabled_client:
        response = enabled_client.post(
            "/api/analyze",
            data={"job_description": "   ", "resume_text": "Vue engineer"},
        )

    assert response.status_code == 400
    with create_engine(database_url).connect() as connection:
        count = connection.execute(
            text("SELECT count(*) FROM analysis_history")
        ).scalar_one()
    assert count == 0


def test_ai_failure_creates_no_history_record(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database_url = f"sqlite:///{tmp_path / 'ai-failure-history.db'}"
    migrate_database(database_url)
    monkeypatch.setattr(settings, "analysis_provider", "unavailable")

    with TestClient(
        create_app(history_enabled=True, database_url=database_url)
    ) as enabled_client:
        response = enabled_client.post(
            "/api/analyze",
            data={
                "job_description": "Build Vue apps.",
                "resume_text": "Vue engineer",
            },
        )

    assert response.status_code == 500
    with create_engine(database_url).connect() as connection:
        count = connection.execute(
            text("SELECT count(*) FROM analysis_history")
        ).scalar_one()
    assert count == 0


def test_persistence_failure_preserves_successful_analysis(tmp_path: Path) -> None:
    database_url = f"sqlite:///{tmp_path / 'broken-history.db'}"
    migrate_database(database_url)
    application = create_app(history_enabled=True, database_url=database_url)

    with TestClient(application) as enabled_client:
        with application.state.history_engine.begin() as connection:
            connection.exec_driver_sql("DROP TABLE analysis_history")
        response = enabled_client.post(
            "/api/analyze",
            data={
                "job_description": "Build Vue apps.",
                "resume_text": "Vue engineer",
            },
        )

    assert response.status_code == 200
    assert response.json()["match_score"] == 78
    assert response.json()["history_id"] is None
    assert response.json()["history_status"] == "failed"


def test_analyze_returns_structured_job_fit_result() -> None:
    response = client.post(
        "/api/analyze",
        data={
            "job_description": "Build and maintain Vue applications.",
            "resume_text": "Frontend engineer with three years of Vue experience.",
        },
    )

    assert response.status_code == 200
    assert response.json()["match_score"] == 78
    assert set(response.json()) == {
        "match_score",
        "summary",
        "strengths",
        "matched_requirements",
        "gaps",
        "recommendations",
        "interview_focus",
        "history_id",
        "history_status",
    }
    assert response.json()["history_id"] is None
    assert response.json()["history_status"] == "disabled"


def test_analyze_rejects_blank_job_description() -> None:
    response = client.post(
        "/api/analyze",
        data={
            "job_description": "   ",
            "resume_text": "Frontend engineer with Vue experience.",
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"]


def test_analyze_rejects_missing_resume() -> None:
    response = client.post(
        "/api/analyze",
        data={"job_description": "Build and maintain Vue applications."},
    )

    assert response.status_code == 400
    assert response.json()["detail"]


def test_analyze_accepts_utf8_text_resume_upload() -> None:
    response = client.post(
        "/api/analyze",
        data={"job_description": "Build and maintain Vue applications."},
        files={
            "resume_file": (
                "resume.txt",
                "Frontend engineer with Vue experience.".encode(),
                "text/plain",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["match_score"] == 78


def test_analyze_rejects_unsupported_resume_file_type() -> None:
    response = client.post(
        "/api/analyze",
        data={"job_description": "Build and maintain Vue applications."},
        files={"resume_file": ("resume.csv", b"Vue,TypeScript", "text/csv")},
    )

    assert response.status_code == 400
    assert response.json()["detail"]


def test_analyze_rejects_resume_larger_than_five_megabytes() -> None:
    response = client.post(
        "/api/analyze",
        data={"job_description": "Build and maintain Vue applications."},
        files={
            "resume_file": (
                "resume.txt",
                b"a" * (5 * 1024 * 1024 + 1),
                "text/plain",
            )
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"]


def test_analyze_rejects_blank_text_resume_upload() -> None:
    response = client.post(
        "/api/analyze",
        data={"job_description": "Build and maintain Vue applications."},
        files={"resume_file": ("resume.txt", b"  \n\t", "text/plain")},
    )

    assert response.status_code == 400
    assert response.json()["detail"]


def test_analyze_rejects_non_utf8_text_resume_upload() -> None:
    response = client.post(
        "/api/analyze",
        data={"job_description": "Build and maintain Vue applications."},
        files={"resume_file": ("resume.txt", b"\xff\xfe\xfa", "text/plain")},
    )

    assert response.status_code == 400
    assert response.json()["detail"]


def test_analyze_accepts_valid_docx_resume_upload() -> None:
    document_bytes = BytesIO()
    document = Document()
    document.add_paragraph("Frontend engineer with Vue experience.")
    document.save(document_bytes)

    response = client.post(
        "/api/analyze",
        data={"job_description": "Build and maintain Vue applications."},
        files={
            "resume_file": (
                "resume.docx",
                document_bytes.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["match_score"] == 78


def test_analyze_accepts_valid_pdf_resume_upload() -> None:
    document_bytes = BytesIO()
    document = canvas.Canvas(document_bytes)
    document.drawString(72, 720, "Frontend engineer with Vue experience.")
    document.save()

    response = client.post(
        "/api/analyze",
        data={"job_description": "Build and maintain Vue applications."},
        files={
            "resume_file": (
                "resume.pdf",
                document_bytes.getvalue(),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["match_score"] == 78
